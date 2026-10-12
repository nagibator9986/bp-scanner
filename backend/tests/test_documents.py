import io

from docx import Document

from tests.conftest import auth


def _division(client, token):
    return client.post("/api/divisions", json={"name": "Казначейство"}, headers=auth(token)).json()["id"]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Регламент обработки платежей")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Шаг"
    table.rows[0].cells[1].text = "Ответственный"
    table.rows[1].cells[0].text = "Проверка реквизитов"
    table.rows[1].cells[1].text = "Операционист"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_upload_extracts_text_including_tables(client, user_token):
    division_id = _division(client, user_token)
    files = {"file": ("Регламент.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post(f"/api/divisions/{division_id}/documents", files=files, headers=auth(user_token))
    assert response.status_code == 200, response.text
    document = response.json()
    assert "Проверка реквизитов | Операционист" in document["text_preview"]
    assert "file_path" not in document and "extracted_text" not in document

    listed = client.get(f"/api/divisions/{division_id}/documents", headers=auth(user_token)).json()
    assert len(listed) == 1 and listed[0]["filename"] == "Регламент.docx"

    download = client.get(f"/api/divisions/{division_id}/documents/{document['id']}/download", headers=auth(user_token))
    assert download.status_code == 200 and download.content[:2] == b"PK"
    assert "filename*=UTF-8''" in download.headers["content-disposition"]

    assert client.delete(f"/api/divisions/{division_id}/documents/{document['id']}", headers=auth(user_token)).status_code == 200
    assert client.get(f"/api/divisions/{division_id}/documents", headers=auth(user_token)).json() == []


def test_upload_cp1251_text(client, user_token):
    division_id = _division(client, user_token)
    files = {"file": ("notes.txt", "Сверка остатков по счетам".encode("cp1251"), "text/plain")}
    document = client.post(f"/api/divisions/{division_id}/documents", files=files, headers=auth(user_token)).json()
    assert document["text_preview"] == "Сверка остатков по счетам"


def test_upload_validation(client, user_token, tmp_settings):
    division_id = _division(client, user_token)
    bad_type = client.post(f"/api/divisions/{division_id}/documents", files={"file": ("virus.exe", b"MZ", "application/octet-stream")}, headers=auth(user_token))
    assert bad_type.status_code == 400

    empty = client.post(f"/api/divisions/{division_id}/documents", files={"file": ("empty.txt", b"", "text/plain")}, headers=auth(user_token))
    assert empty.status_code == 400

    tmp_settings(max_upload_mb=0)
    too_big = client.post(f"/api/divisions/{division_id}/documents", files={"file": ("big.txt", b"x" * 10, "text/plain")}, headers=auth(user_token))
    assert too_big.status_code == 413


def test_extract_text_endpoint(client, user_token):
    response = client.post("/api/extract-text", files={"file": ("a.txt", "Привет".encode(), "text/plain")}, headers=auth(user_token))
    assert response.json()["text"] == "Привет"


def test_questionnaire_completion(client, user_token):
    division_id = _division(client, user_token)
    answers = {"dept_function": "Платежи", "employees_roles": "Операционисты", "has_approvals": False}
    saved = client.post(f"/api/divisions/{division_id}/questionnaire", json={"answers": answers}, headers=auth(user_token)).json()
    assert saved["completion"]["required_done"] == 2
    loaded = client.get(f"/api/divisions/{division_id}/questionnaire", headers=auth(user_token)).json()
    assert loaded["answers"]["dept_function"] == "Платежи"
    division = client.get(f"/api/divisions/{division_id}", headers=auth(user_token)).json()
    assert division["questionnaire_completion"]["required_done"] == 2
