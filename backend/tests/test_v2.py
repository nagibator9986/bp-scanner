"""Конспект документов, экономика, редактирование процессов, инициативы, частичный перезапуск, Excel."""
import io

from openpyxl import load_workbook

from bpscanner.context import split_text
from bpscanner.economics import DEFAULT_ECONOMICS, case_minutes, frequency_to_monthly, summarize, tobe_economics
from bpscanner.normalize import normalize_tobe
from bpscanner.pipeline import dedupe_tobe
from tests.conftest import auth, wait_for_analysis
from tests.test_pipeline import _prepare_division


# ---------- TO-BE пакетами ----------

def test_dedupe_tobe_across_batches():
    processes = [{"id": f"p{i}", "name": name, "steps_detail": [{"name": "a"}, {"name": "b"}]}
                 for i, name in enumerate(["Приём заявок", "Согласование договора", "Архивирование дел"])]
    tobe = lambda name: {"process_name": name, "new_process": {"steps": [{"name": "Шаг"}]}}  # noqa: E731
    # Пакет 1 — первые два процесса; пакет 2 — третий, но модель снова описала «Приём заявок» и «Согласование»
    models = normalize_tobe([tobe("Приём заявок"), tobe("Нечто новое"), tobe("Совсем другое")], processes[:2])
    models += normalize_tobe([tobe("Приём заявок"), tobe("Согласование договора"), tobe("Архивирование дел")], processes[2:])
    result = dedupe_tobe(models, processes)
    by_process = {m["process_id"]: m for m in result}
    assert set(by_process) == {"p0", "p1", "p2", None}
    assert sum(1 for m in result if m["process_id"] == "p0") == 1
    # «Согласование договора» не имело модели в своём пакете — описание из чужого пакета привязано к нему
    assert by_process["p1"]["process_name"] == "Согласование договора"
    assert by_process["p1"]["metrics"]["steps_before"] == 2


# ---------- Экономика (юнит) ----------

def test_frequency_to_monthly():
    assert frequency_to_monthly("Ежедневно") == 21
    assert frequency_to_monthly("2 раза в неделю") == 8.6
    assert round(frequency_to_monthly("Ежеквартально"), 2) == 0.33
    assert frequency_to_monthly("по запросу") is None


def test_case_minutes_and_tobe_economics():
    steps = [
        {"name": "Принять", "duration_min": 10},
        {"name": "Ок?", "type": "decision"},
        {"name": "Ввести", "duration_min": 20},
        {"name": "Отправить", "duration_min": 5, "automated": True},
    ]
    assert case_minutes(steps) == 30
    assert case_minutes([{"name": "a"}]) is None
    process = {"id": "p", "name": "P", "volume_per_month": 120, "steps_detail": steps}
    tobe = {"new_process": {"steps_detail": [
        {"name": "Принять", "duration_min": None},
        {"name": "Ввести автоматически", "automated": True},
        {"name": "Проверить", "duration_min": 3},
    ]}}
    result = tobe_economics(tobe, process)
    assert result["minutes_per_case"] == 13  # 10 (из AS-IS) + 0 + 3
    assert result["hours_per_month"] == 26 and result["saved_hours_month"] == 34
    assert result["reduction_percent"] == 57


def test_summarize_money_and_fte():
    process = {"id": "p", "name": "P", "economics": {"volume_per_month": 100, "minutes_per_case": 66, "hours_per_month": 110}}
    tobe = {"process_id": "p", "economics": {"minutes_per_case": 33, "hours_per_month": 55, "saved_hours_month": 55}}
    eco = {**DEFAULT_ECONOMICS, "hourly_cost": 1000, "fte_hours_month": 110}
    result = summarize([process], [tobe], eco)
    assert result["totals"]["fte"] == 1.0 and result["totals"]["saved_fte"] == 0.5
    assert result["totals"]["saved_money_year"] == 55 * 12 * 1000
    assert summarize([process], [tobe], DEFAULT_ECONOMICS)["totals"]["saved_money_year"] is None


def test_split_text_respects_size():
    text = "\n".join(f"Абзац {i} " + "слово " * 30 for i in range(50))
    parts = split_text(text, 1000)
    assert all(len(p) <= 1000 for p in parts)
    assert "".join(p.replace("\n", "") for p in parts) == text.replace("\n", "")


# ---------- Интеграция ----------

def _run(client, token, division_id, mode="full"):
    response = client.post(f"/api/divisions/{division_id}/analyze", json={"mode": mode}, headers=auth(token))
    assert response.status_code == 200, response.text
    status = wait_for_analysis(client, token, division_id)
    assert status["status"] == "completed", status
    return status


def test_economics_runs_and_rescoring_of_missing(client, admin_token, user_token, configure_llm, fake_llm):
    configure_llm()
    division_id = _prepare_division(client, user_token)
    _run(client, user_token, division_id)

    # Модель «забыла» второй процесс — он оценён дозапросом
    scoring = client.get(f"/api/divisions/{division_id}/scoring", headers=auth(user_token)).json()
    assert len(scoring) == 2
    assert sum("Оцени процессы" in p for p in fake_llm.user_prompts()) == 2

    economics = client.get(f"/api/divisions/{division_id}/economics", headers=auth(user_token)).json()
    row = next(r for r in economics["rows"] if r["process_name"] == "Приём и регистрация заявок")
    assert row["minutes_per_case"] == 24 and row["hours_per_month"] == 240
    assert row["tobe_minutes_per_case"] == 9 and row["saved_hours_month"] == 150
    assert economics["totals"]["saved_money_year"] is None  # стоимость часа не задана

    assert client.post("/api/settings/economics", json={"hourly_cost": 5000}, headers=auth(user_token)).status_code == 403
    client.post("/api/settings/economics", json={"hourly_cost": 5000, "currency": "₸", "fte_hours_month": 150}, headers=auth(admin_token))
    economics = client.get(f"/api/divisions/{division_id}/economics", headers=auth(user_token)).json()
    assert economics["totals"]["saved_money_year"] == 150 * 12 * 5000
    assert economics["totals"]["saved_fte"] == 1.0

    dashboard = client.get("/api/dashboard", headers=auth(user_token)).json()
    assert dashboard["economics"]["saved_hours_month"] == 150
    assert dashboard["divisions"][0]["saved_money_year"] == 9_000_000
    assert sum(s["count"] for s in dashboard["initiatives_by_status"]) == 3

    runs = client.get(f"/api/divisions/{division_id}/runs", headers=auth(user_token)).json()
    assert runs[0]["status"] == "completed" and runs[0]["model"] == "qwen-test" and runs[0]["mode"] == "full"
    assert runs[0]["summary"]["processes_count"] == 2

    pdf = client.get(f"/api/divisions/{division_id}/export/pdf", headers=auth(user_token))
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    word = client.get(f"/api/divisions/{division_id}/export/word", headers=auth(user_token))
    assert word.status_code == 200


def test_long_documents_are_digested_and_cached(client, user_token, configure_llm, fake_llm, tmp_settings):
    configure_llm()
    tmp_settings(max_document_chars=3000, digest_chunk_chars=2000)
    division_id = _prepare_division(client, user_token)
    long_text = "\n".join(f"Пункт {i}. Операционист регистрирует заявку и проверяет комплект документов." for i in range(150))
    client.post(f"/api/divisions/{division_id}/documents", files={"file": ("регламент_большой.txt", long_text.encode(), "text/plain")},
                headers=auth(user_token))

    _run(client, user_token, division_id)
    digests = [p for p in fake_llm.user_prompts() if "Фрагмент документа" in p]
    assert len(digests) >= 4
    extraction = next(p for p in fake_llm.user_prompts() if "Выяви ключевые бизнес-процессы" in p)
    assert "из конспекта документа" in extraction

    fake_llm.requests.clear()
    _run(client, user_token, division_id)
    assert not [p for p in fake_llm.user_prompts() if "Фрагмент документа" in p], "конспекты должны браться из кэша"


def test_process_editing_partial_rerun_and_initiatives(client, user_token, configure_llm, fake_llm):
    configure_llm()
    division_id = _prepare_division(client, user_token)
    _run(client, user_token, division_id)
    headers = auth(user_token)

    # Инициатива: статус, ответственный, комментарий
    recs = client.get(f"/api/divisions/{division_id}/recommendations", headers=headers).json()
    target = recs[0]
    assert target["status"] == "new"
    updated = client.patch(f"/api/divisions/{division_id}/recommendations/{target['id']}",
                           json={"status": "accepted", "owner": "Петров", "comment": "Старт в Q1", "due_date": "2027-03-01"},
                           headers=headers).json()
    assert updated["status"] == "accepted" and updated["status_updated_at"]
    bad = client.patch(f"/api/divisions/{division_id}/recommendations/{target['id']}", json={"status": "unknown"}, headers=headers)
    assert bad.status_code == 422

    # Редактирование процесса перестраивает BPMN и помечает результаты устаревшими
    process = client.get(f"/api/divisions/{division_id}/processes", headers=headers).json()[1]
    payload = {
        "name": "Подготовка отчётности (уточнено)",
        "frequency": "Ежемесячно",
        "volume_per_month": 2,
        "steps": [
            {"name": "Выгрузить данные", "role": "Аналитик", "system": "АБИС", "duration_min": 60},
            {"name": "Данные сходятся?", "role": "Аналитик", "type": "decision"},
            {"name": "Отправить отчёт", "role": "Руководитель", "duration_min": 15},
        ],
    }
    edited = client.put(f"/api/divisions/{division_id}/processes/{process['id']}", json=payload, headers=headers).json()
    assert edited["edited"] and edited["roles"] == ["Аналитик", "Руководитель"]
    assert edited["economics"]["hours_per_month"] == 2.5
    bpmn = client.get(f"/api/divisions/{division_id}/bpmn?type=as_is", headers=headers).json()
    model = next(m for m in bpmn if m["process_id"] == process["id"])
    assert "Данные сходятся?" in model["xml"] and model["process_name"] == payload["name"]
    division = client.get(f"/api/divisions/{division_id}", headers=headers).json()
    assert division["results_outdated"] is True

    created = client.post(f"/api/divisions/{division_id}/processes", json={"name": "Новый процесс", "steps": [{"name": "Шаг"}]},
                          headers=headers).json()
    assert client.get(f"/api/divisions/{division_id}", headers=headers).json()["processes_count"] == 3

    # Пересчёт скоринга и рекомендаций по отредактированным процессам
    fake_llm.requests.clear()
    status = _run(client, user_token, division_id, mode="scoring")
    assert {s["key"]: s["state"] for s in status["steps"]}["extracting_processes"] == "reused"
    assert not any("Выяви ключевые бизнес-процессы" in p for p in fake_llm.user_prompts())
    assert any("Подготовка отчётности (уточнено)" in p for p in fake_llm.user_prompts())
    processes = client.get(f"/api/divisions/{division_id}/processes", headers=headers).json()
    assert [p["name"] for p in processes][1] == payload["name"]
    scoring = client.get(f"/api/divisions/{division_id}/scoring", headers=headers).json()
    assert {s["process_id"] for s in scoring} == {p["id"] for p in processes}
    assert client.get(f"/api/divisions/{division_id}", headers=headers).json()["results_outdated"] is False
    # AS-IS схема отредактированного процесса сохранилась
    assert any(m["process_id"] == process["id"] for m in client.get(f"/api/divisions/{division_id}/bpmn?type=as_is", headers=headers).json())

    # Статус инициативы перенесён на перегенерированную рекомендацию
    _run(client, user_token, division_id, mode="recommendations")
    recs = client.get(f"/api/divisions/{division_id}/recommendations", headers=headers).json()
    carried = next(r for r in recs if r["title"] == target["title"])
    assert carried["status"] == "accepted" and carried["owner"] == "Петров" and carried["id"] != target["id"]

    assert client.delete(f"/api/divisions/{division_id}/processes/{created['id']}", headers=headers).status_code == 200
    assert len(client.get(f"/api/divisions/{division_id}/processes", headers=headers).json()) == 2

    runs = client.get(f"/api/divisions/{division_id}/runs", headers=headers).json()
    assert [r["mode"] for r in runs] == ["recommendations", "scoring", "full"]


def test_partial_rerun_requires_processes(client, user_token, configure_llm):
    configure_llm()
    division_id = _prepare_division(client, user_token)
    response = client.post(f"/api/divisions/{division_id}/analyze", json={"mode": "scoring"}, headers=auth(user_token))
    assert response.status_code == 400


def test_questionnaire_suggestions(client, user_token, configure_llm):
    configure_llm()
    division_id = client.post("/api/divisions", json={"name": "Без документов"}, headers=auth(user_token)).json()["id"]
    assert client.post(f"/api/divisions/{division_id}/questionnaire/suggest", headers=auth(user_token)).status_code == 400

    division_id = _prepare_division(client, user_token)
    result = client.post(f"/api/divisions/{division_id}/questionnaire/suggest", headers=auth(user_token)).json()
    assert result["suggestions"] == {
        "dept_function": "Обслуживание клиентов МСБ",
        "has_approvals": True,
        "process_frequency": "Ежедневно",
    }


def test_excel_exports(client, user_token, configure_llm):
    configure_llm()
    division_id = _prepare_division(client, user_token)
    _run(client, user_token, division_id)

    response = client.get(f"/api/divisions/{division_id}/export/xlsx", headers=auth(user_token))
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Сводка", "Анкета", "Процессы", "Шаги", "Скоринг", "Экономика", "Инициативы", "TO-BE", "n8n"]
    assert workbook["Процессы"].max_row == 3
    assert workbook["Инициативы"]["C2"].value == "Новая"

    portfolio = client.get("/api/dashboard/export/xlsx", headers=auth(user_token))
    workbook = load_workbook(io.BytesIO(portfolio.content))
    assert workbook.sheetnames == ["Подразделения", "Инициативы", "Процессы"]
    assert workbook["Инициативы"].max_row == 4
