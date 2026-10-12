"""Выгрузка результатов в Excel: по подразделению и по всему портфелю."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import SCORE_LABELS
from .normalize import CATEGORY_LABELS, INITIATIVE_STATUSES, SCORE_FIELDS
from .questionnaire import QUESTION_LABELS

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
COMPLEXITY_RU = {"low": "низкая", "medium": "средняя", "high": "высокая"}


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    sheet = workbook.create_sheet(title[:31])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(["; ".join(map(str, v)) if isinstance(v, list) else v for v in row])
    for index, header in enumerate(headers, 1):
        width = (widths or [])[index - 1] if widths and index <= len(widths) else max(12, min(50, len(header) + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _save(workbook: Workbook) -> bytes:
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
        del workbook["Sheet"]
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def division_workbook(data: dict, economics: dict) -> bytes:
    division = data["division"]
    eco = economics["settings"]
    totals = economics["totals"]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    rows = [
        ("Подразделение", division.get("name")),
        ("Руководитель", division.get("head_name")),
        ("Дата анализа", (division.get("analysis_completed_at") or "")[:10]),
        ("Процессов", len(data["processes"])),
        ("Средний потенциал автоматизации, %", division.get("avg_automation_potential")),
        ("Рекомендаций", len(data["recommendations"])),
        ("Трудозатраты AS-IS, ч/мес", totals.get("hours_per_month")),
        ("Высвобождается, ч/мес", totals.get("saved_hours_month")),
        ("Высвобождается, FTE", totals.get("saved_fte")),
        (f"Экономия в год, {eco['currency']}", totals.get("saved_money_year")),
        ("Стоимость часа", eco.get("hourly_cost") or "не задана"),
        ("Рабочих часов в месяц на FTE", eco.get("fte_hours_month")),
    ]
    for label, value in rows:
        summary.append([label, value])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 40

    answers = (data.get("questionnaire") or {}).get("answers") or {}
    _sheet(workbook, "Анкета", ["Вопрос", "Ответ"],
           [[QUESTION_LABELS.get(k, k), ("Да" if v else "Нет") if isinstance(v, bool) else v] for k, v in answers.items()], [50, 90])

    scores = {s.get("process_id"): s for s in data["scoring"]}
    _sheet(
        workbook, "Процессы",
        ["Процесс", "Описание", "Владелец", "Частота", "Экз./мес", "SLA", "Шагов", "Решений", "Роли", "Системы", "Потенциал, %", "Проблемы"],
        [[p.get("name"), p.get("description"), p.get("owner"), p.get("frequency"), (p.get("economics") or {}).get("volume_per_month"),
          p.get("sla"), (p.get("metrics") or {}).get("steps_count"), (p.get("metrics") or {}).get("decisions_count"),
          p.get("roles"), p.get("systems"), (scores.get(p.get("id")) or {}).get("automation_potential"), p.get("pain_points")]
         for p in data["processes"]],
        [36, 50, 24, 18, 10, 16, 8, 9, 30, 30, 12, 40],
    )
    step_rows = []
    for process in data["processes"]:
        for index, step in enumerate(process.get("steps_detail") or [], 1):
            step_rows.append([process.get("name"), index, step.get("name"), step.get("role"), step.get("system"),
                              "решение" if step.get("type") == "decision" else "задача",
                              "авто" if step.get("automated") else ("вручную" if step.get("manual") else ""), step.get("duration_min")])
    _sheet(workbook, "Шаги", ["Процесс", "№", "Шаг", "Роль", "Система", "Тип", "Выполнение", "Минут"], step_rows, [32, 5, 44, 24, 20, 10, 12, 8])

    _sheet(
        workbook, "Скоринг",
        ["Процесс"] + [SCORE_LABELS[f] for f in SCORE_FIELDS] + ["Потенциал, %", "Узкие места", "Обоснование"],
        [[s.get("process_name")] + [s.get(f) for f in SCORE_FIELDS] + [s.get("automation_potential"), s.get("bottlenecks"), s.get("rationale")]
         for s in data["scoring"]],
        [36, 11, 11, 13, 13, 13, 14, 12, 40, 50],
    )
    _sheet(
        workbook, "Экономика",
        ["Процесс", "Экз./мес", "Мин. на экз. AS-IS", "Ч/мес AS-IS", "FTE AS-IS", "Мин. на экз. TO-BE", "Ч/мес TO-BE",
         "Высвобождается ч/мес", "Высвобождается FTE", f"Экономия в год, {eco['currency']}", "Сокращение, %"],
        [[r["process_name"], r["volume_per_month"], r["minutes_per_case"], r["hours_per_month"], r["fte"], r["tobe_minutes_per_case"],
          r["tobe_hours_per_month"], r["saved_hours_month"], r["saved_fte"], r["saved_money_year"], r["reduction_percent"]]
         for r in economics["rows"]],
        [36, 10, 14, 12, 10, 14, 12, 14, 14, 18, 12],
    )
    _sheet(
        workbook, "Инициативы",
        ["Рекомендация", "Категория", "Статус", "Ответственный", "Срок", "Приоритет", "Сложность", "FTE (оценка модели)",
         "Недель", "Эффект", "Стек", "Процессы", "Комментарий"],
        [[r.get("title"), CATEGORY_LABELS.get(r.get("category"), r.get("category")), INITIATIVE_STATUSES.get(r.get("status", "new")),
          r.get("owner"), r.get("due_date"), r.get("priority"), COMPLEXITY_RU.get(r.get("complexity")), r.get("estimated_fte_savings"),
          r.get("implementation_weeks"), r.get("effect"), r.get("implementation_stack"), r.get("related_processes"), r.get("comment")]
         for r in data["recommendations"]],
        [40, 24, 12, 22, 12, 10, 11, 12, 8, 40, 24, 36, 36],
    )
    _sheet(
        workbook, "TO-BE",
        ["Процесс", "Целевой процесс", "Шаги TO-BE", "Улучшения", "Автоматизируется", "Исключается", "Сокращение, %", "Эффект", "Риски"],
        [[t.get("process_name"), (t.get("new_process") or {}).get("name"), (t.get("new_process") or {}).get("steps"), t.get("improvements"),
          t.get("automated_steps"), t.get("removed_steps"), (t.get("economics") or {}).get("reduction_percent") or t.get("time_reduction_percent"),
          t.get("expected_effect"), t.get("risks")]
         for t in data["tobe"]],
        [32, 32, 50, 40, 32, 32, 12, 36, 32],
    )
    _sheet(
        workbook, "n8n",
        ["Workflow", "Описание", "Экономия времени", "Интеграции", "Ноды"],
        [[w.get("name"), w.get("description"), w.get("estimated_time_savings"), w.get("integration_points"),
          " → ".join(n.get("name", "") for n in w.get("nodes") or [])] for w in data["n8n_workflows"]],
        [36, 50, 18, 30, 70],
    )
    return _save(workbook)


def portfolio_workbook(dashboard: dict, recommendations: list[dict], processes: list[dict], scoring: list[dict]) -> bytes:
    names = {d["id"]: d.get("name", "") for d in dashboard["divisions"]}
    currency = dashboard.get("economics", {}).get("currency", "₸")
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(
        workbook, "Подразделения",
        ["Подразделение", "Руководитель", "Статус анализа", "Заполненность анкеты, %", "Документов", "Процессов",
         "Потенциал, %", "Рекомендаций", "Высвобождается ч/мес", f"Экономия в год, {currency}", "Дата анализа"],
        [[d.get("name"), d.get("head_name"), d.get("analysis_status"), (d.get("questionnaire_completion") or {}).get("percent"),
          d.get("documents_count"), d.get("processes_count"), d.get("avg_automation_potential"), d.get("recommendations_count"),
          d.get("saved_hours_month"), d.get("saved_money_year"), (d.get("analysis_completed_at") or "")[:10]]
         for d in dashboard["divisions"]],
        [36, 24, 14, 14, 11, 11, 12, 13, 16, 18, 14],
    )
    _sheet(
        workbook, "Инициативы",
        ["Подразделение", "Рекомендация", "Категория", "Статус", "Ответственный", "Срок", "Приоритет", "Сложность", "FTE", "Недель", "Комментарий"],
        [[names.get(r.get("division_id")), r.get("title"), CATEGORY_LABELS.get(r.get("category"), r.get("category")),
          INITIATIVE_STATUSES.get(r.get("status", "new")), r.get("owner"), r.get("due_date"), r.get("priority"),
          COMPLEXITY_RU.get(r.get("complexity")), r.get("estimated_fte_savings"), r.get("implementation_weeks"), r.get("comment")]
         for r in recommendations],
        [30, 44, 24, 12, 22, 12, 10, 11, 8, 8, 36],
    )
    scores = {s.get("process_id"): s for s in scoring}
    _sheet(
        workbook, "Процессы",
        ["Подразделение", "Процесс", "Экз./мес", "Ч/мес AS-IS", "Потенциал, %"] + [SCORE_LABELS[f] for f in SCORE_FIELDS],
        [[names.get(p.get("division_id")), p.get("name"), (p.get("economics") or {}).get("volume_per_month"),
          (p.get("economics") or {}).get("hours_per_month"), (scores.get(p.get("id")) or {}).get("automation_potential")]
         + [(scores.get(p.get("id")) or {}).get(f) for f in SCORE_FIELDS]
         for p in processes],
        [30, 40, 10, 12, 12, 11, 11, 13, 13, 13, 14],
    )
    return _save(workbook)
