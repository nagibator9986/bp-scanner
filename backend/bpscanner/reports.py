"""Отчёты по анализу подразделения: Word (.docx) и PDF (с диаграммами).

Генерация синхронная и тяжёлая — вызывайте через asyncio.to_thread.
"""
import math
from datetime import datetime
from html import escape
from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .analytics import SCORE_LABELS
from .normalize import CATEGORY_LABELS, SCORE_FIELDS
from .questionnaire import QUESTION_LABELS

COMPLEXITY_RU = {"low": "низкая", "medium": "средняя", "high": "высокая"}
ACCENT = "1E3A5F"


def _date(value) -> str:
    if not value:
        return "—"
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).strftime("%d.%m.%Y")
    except ValueError:
        return str(value)[:10]


def _answer(value) -> str:
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, list):
        return ", ".join(map(str, value))
    return str(value) if value not in (None, "") else "—"


def _score_color(value: float, inverse: bool = False) -> str:
    """Цвет ячейки скоринга: зелёный → жёлтый → красный."""
    value = max(1, min(10, float(value or 1)))
    ratio = (value - 1) / 9
    if inverse:
        ratio = 1 - ratio
    if ratio < 0.5:
        r, g = int(220 + 35 * ratio * 2), 235
    else:
        r, g = 255, int(235 - 120 * (ratio - 0.5) * 2)
    return f"{r:02X}{g:02X}B0"


def summary(data: dict) -> dict:
    recommendations = data["recommendations"]
    scoring = data["scoring"]
    potentials = [s.get("automation_potential") for s in scoring if isinstance(s.get("automation_potential"), (int, float))]
    return {
        "processes": len(data["processes"]),
        "recommendations": len(recommendations),
        "workflows": len(data["n8n_workflows"]),
        "tobe": len(data["tobe"]),
        "fte": round(sum(r.get("estimated_fte_savings") or 0 for r in recommendations), 1),
        "potential": round(sum(potentials) / len(potentials)) if potentials else None,
    }


def content_disposition(filename: str, fallback: str) -> str:
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quote(filename)}"


def report_filename(division: dict, extension: str) -> tuple[str, str]:
    name = "".join(ch for ch in division.get("name", "report") if ch.isalnum() or ch in " -_").strip() or "report"
    stamp = datetime.now().strftime("%Y-%m-%d")
    return f"Анализ_{name.replace(' ', '_')}_{stamp}.{extension}", f"bp_report_{stamp}.{extension}"


# ============== Word ==============

def _shade(cell, hex_color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    properties.append(shading)


def _table(document, headers: list[str], rows: list[list], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, ACCENT)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
    return table


def _bullets(document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def _label_value(document, label: str, value) -> None:
    if value in (None, "", []):
        return
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(", ".join(value) if isinstance(value, list) else str(value))


def build_word_report(data: dict) -> bytes:
    division = data["division"]
    stats = summary(data)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = document.add_heading("Отчёт по анализу бизнес-процессов", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _label_value(document, "Подразделение", division.get("name"))
    _label_value(document, "Руководитель", division.get("head_name"))
    _label_value(document, "Дата анализа", _date(division.get("analysis_completed_at")))

    document.add_heading("Ключевые показатели", level=1)
    _table(document, ["Показатель", "Значение"], [
        ["Выявлено процессов", stats["processes"]],
        ["Средний потенциал автоматизации", f"{stats['potential']}%" if stats["potential"] is not None else "—"],
        ["Рекомендаций", stats["recommendations"]],
        ["Потенциальная экономия, FTE", stats["fte"]],
        ["TO-BE моделей", stats["tobe"]],
        ["n8n workflows", stats["workflows"]],
    ])
    for warning in division.get("analysis_warnings") or []:
        document.add_paragraph(f"⚠ {warning}")

    document.add_heading("1. Анкета подразделения", level=1)
    answers = (data.get("questionnaire") or {}).get("answers") or {}
    if answers:
        _table(document, ["Вопрос", "Ответ"], [[QUESTION_LABELS.get(k, k), _answer(v)] for k, v in answers.items()])
    else:
        document.add_paragraph("Анкета не заполнена")

    document.add_heading("2. Выявленные бизнес-процессы", level=1)
    if not data["processes"]:
        document.add_paragraph("Процессы не выявлены")
    for index, process in enumerate(data["processes"], 1):
        document.add_heading(f"2.{index}. {process.get('name', 'Процесс')}", level=2)
        if process.get("description"):
            document.add_paragraph(process["description"])
        _label_value(document, "Владелец", process.get("owner"))
        _label_value(document, "Триггер", process.get("trigger"))
        _label_value(document, "Частота", process.get("frequency"))
        _label_value(document, "SLA", process.get("sla"))
        _label_value(document, "Роли", process.get("roles"))
        _label_value(document, "Системы", process.get("systems"))
        steps = process.get("steps_detail") or [{"name": s} for s in process.get("steps") or []]
        if steps:
            document.add_paragraph().add_run("Шаги процесса:").bold = True
            for number, step in enumerate(steps, 1):
                extra = ", ".join(x for x in (step.get("role"), step.get("system")) if x)
                marker = "◇ " if step.get("type") == "decision" else ""
                document.add_paragraph(f"{marker}{step['name']}" + (f" ({extra})" if extra else ""), style="List Number")
        if process.get("pain_points"):
            document.add_paragraph().add_run("Проблемы:").bold = True
            _bullets(document, process["pain_points"])

    document.add_heading("3. Скоринг процессов", level=1)
    if data["scoring"]:
        headers = ["Процесс"] + [SCORE_LABELS[f] for f in SCORE_FIELDS] + ["Потенциал"]
        table = _table(document, headers, [
            [s.get("process_name", "")] + [s.get(f, "") for f in SCORE_FIELDS] + [f"{s.get('automation_potential', '—')}%"]
            for s in data["scoring"]
        ])
        for row_index, score in enumerate(data["scoring"], 1):
            for column, field in enumerate(SCORE_FIELDS, 1):
                if isinstance(score.get(field), (int, float)):
                    _shade(table.rows[row_index].cells[column], _score_color(score[field]))
        document.add_paragraph("Шкала 1–10. Потенциал автоматизации — интегральный индекс 0–100%.").runs[0].italic = True
    else:
        document.add_paragraph("Скоринг не выполнен")

    document.add_heading("4. Рекомендации", level=1)
    if not data["recommendations"]:
        document.add_paragraph("Рекомендации не сформированы")
    for number, (key, label) in enumerate(CATEGORY_LABELS.items(), 1):
        items = [r for r in data["recommendations"] if r.get("category") == key]
        if not items:
            continue
        document.add_heading(f"4.{number}. {label}", level=2)
        _table(document, ["Рекомендация", "Приоритет", "Сложность", "FTE", "Срок, нед."], [
            [r.get("title", ""), r.get("priority", "—"), COMPLEXITY_RU.get(r.get("complexity"), "—"),
             r.get("estimated_fte_savings", 0), r.get("implementation_weeks") or "—"]
            for r in items
        ])
        for recommendation in items:
            paragraph = document.add_paragraph()
            paragraph.add_run(recommendation.get("title", "")).bold = True
            if recommendation.get("description"):
                document.add_paragraph(recommendation["description"])
            _label_value(document, "Эффект", recommendation.get("effect"))
            _label_value(document, "Стек", recommendation.get("implementation_stack"))
            _label_value(document, "Процессы", recommendation.get("related_processes"))

    document.add_heading("5. TO-BE модели", level=1)
    if not data["tobe"]:
        document.add_paragraph("TO-BE модели не сформированы")
    for index, model in enumerate(data["tobe"], 1):
        document.add_heading(f"5.{index}. {model.get('process_name', '')}", level=2)
        metrics = model.get("metrics") or {}
        if metrics.get("steps_before") is not None:
            document.add_paragraph(
                f"Шагов: {metrics.get('steps_before')} → {metrics.get('steps_after')}, "
                f"автоматизировано: {metrics.get('automated_count', 0)}"
                + (f", сокращение времени ≈ {model['time_reduction_percent']}%" if model.get("time_reduction_percent") else "")
            )
        if model.get("improvements"):
            document.add_paragraph().add_run("Улучшения:").bold = True
            _bullets(document, model["improvements"])
        new_steps = (model.get("new_process") or {}).get("steps_detail") or []
        if new_steps:
            document.add_paragraph().add_run("Целевой процесс:").bold = True
            for step in new_steps:
                suffix = " [авто]" if step.get("automated") else ""
                document.add_paragraph(f"{step['name']}{suffix}", style="List Number")
        _label_value(document, "Ожидаемый эффект", model.get("expected_effect"))
        _label_value(document, "Риски", model.get("risks"))

    document.add_heading("6. n8n workflows", level=1)
    if not data["n8n_workflows"]:
        document.add_paragraph("Workflows не сгенерированы")
    for index, workflow in enumerate(data["n8n_workflows"], 1):
        document.add_heading(f"6.{index}. {workflow.get('name', 'Workflow')}", level=2)
        if workflow.get("description"):
            document.add_paragraph(workflow["description"])
        _label_value(document, "Экономия времени", workflow.get("estimated_time_savings"))
        _label_value(document, "Интеграции", workflow.get("integration_points"))
        nodes = workflow.get("nodes") or []
        if nodes:
            document.add_paragraph("Цепочка: " + " → ".join(n.get("name", "") for n in nodes))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ============== PDF ==============

def _svg_hbar(items: list[tuple[str, float]], max_value: float, unit: str = "", color: str = "#3b82f6") -> str:
    if not items:
        return ""
    row_h, label_w, bar_w = 26, 230, 330
    height = row_h * len(items) + 10
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="640" height="{height}" viewBox="0 0 640 {height}">']
    for index, (label, value) in enumerate(items):
        y = 5 + index * row_h
        width = 0 if max_value <= 0 else max(2, bar_w * value / max_value)
        short = label if len(label) <= 38 else label[:37] + "…"
        parts.append(f'<text x="{label_w - 8}" y="{y + 16}" font-size="10" text-anchor="end" fill="#334155">{escape(short)}</text>')
        parts.append(f'<rect x="{label_w}" y="{y + 4}" width="{bar_w}" height="16" rx="3" fill="#eef2f7"/>')
        parts.append(f'<rect x="{label_w}" y="{y + 4}" width="{width:.1f}" height="16" rx="3" fill="{color}"/>')
        parts.append(f'<text x="{label_w + bar_w + 8}" y="{y + 16}" font-size="10" fill="#0f172a">{value:g}{unit}</text>')
    parts.append("</svg>")
    return "".join(parts)


def _svg_donut(segments: list[tuple[str, float, str]]) -> str:
    total = sum(v for _, v, _ in segments)
    if total <= 0:
        return ""
    cx, cy, r, stroke = 80, 80, 56, 26
    circumference = 2 * math.pi * r
    offset = 0.0
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" width="460" height="160" viewBox="0 0 460 160">']
    for label, value, color in segments:
        length = circumference * value / total
        parts.append(
            f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="{color}" stroke-width="{stroke}" '
            f'stroke-dasharray="{length:.2f} {circumference - length:.2f}" stroke-dashoffset="{-offset:.2f}" '
            f'transform="rotate(-90 {cx} {cy})"/>'
        )
        offset += length
    parts.append(f'<text x="{cx}" y="{cy + 6}" font-size="18" font-weight="bold" text-anchor="middle" fill="#0f172a">{int(total)}</text>')
    for index, (label, value, color) in enumerate(segments):
        y = 40 + index * 28
        parts.append(f'<rect x="180" y="{y - 11}" width="14" height="14" rx="3" fill="{color}"/>')
        parts.append(f'<text x="202" y="{y}" font-size="11" fill="#334155">{escape(label)} — {int(value)}</text>')
    parts.append("</svg>")
    return "".join(parts)


PDF_CSS = """
@page { size: A4; margin: 18mm 15mm 20mm 15mm;
        @bottom-right { content: "Стр. " counter(page) " из " counter(pages); font-size: 8pt; color: #64748b; } }
body { font-family: 'DejaVu Sans', Arial, sans-serif; color: #1e293b; font-size: 9.5pt; line-height: 1.45; }
h1 { color: #1e3a5f; font-size: 20pt; margin: 0 0 4px; }
h2 { color: #1e3a5f; font-size: 13.5pt; border-bottom: 2px solid #3b82f6; padding-bottom: 4px; margin-top: 22px; page-break-after: avoid; }
h3 { color: #334155; font-size: 11pt; margin: 14px 0 6px; page-break-after: avoid; }
.meta { color: #64748b; margin-bottom: 14px; }
.kpis { display: flex; flex-wrap: wrap; gap: 8px; margin: 10px 0; }
.kpi { flex: 1 0 28%; border: 1px solid #e2e8f0; border-radius: 8px; padding: 8px 10px; background: #f8fafc; }
.kpi b { display: block; font-size: 16pt; color: #1d4ed8; }
.kpi span { color: #64748b; font-size: 8pt; }
table { width: 100%; border-collapse: collapse; margin: 8px 0; font-size: 8.5pt; }
th { background: #1e3a5f; color: #fff; text-align: left; padding: 5px 6px; }
td { border: 1px solid #e2e8f0; padding: 5px 6px; vertical-align: top; }
td.num { text-align: center; }
.card { border: 1px solid #e2e8f0; border-left: 3px solid #3b82f6; border-radius: 6px; padding: 8px 10px; margin: 8px 0; page-break-inside: avoid; }
.muted { color: #64748b; }
.tag { display: inline-block; background: #eff6ff; color: #1d4ed8; border-radius: 4px; padding: 1px 6px; margin: 1px 2px; font-size: 8pt; }
.warn { background: #fffbeb; border: 1px solid #fcd34d; border-radius: 6px; padding: 6px 10px; margin: 6px 0; }
ol, ul { margin: 4px 0 4px 18px; padding: 0; }
.chain { color: #475569; font-size: 8.5pt; }
"""


def build_pdf_report(data: dict) -> bytes:
    from weasyprint import HTML
    from weasyprint.urls import URLFetcher

    division = data["division"]
    stats = summary(data)
    e = lambda value: escape(str(value if value is not None else ""))  # noqa: E731
    html = [f"<html><head><meta charset='utf-8'><style>{PDF_CSS}</style></head><body>"]
    html.append("<h1>Отчёт по анализу бизнес-процессов</h1>")
    html.append(
        f"<div class='meta'>Подразделение: <b>{e(division.get('name'))}</b>"
        + (f" · Руководитель: {e(division.get('head_name'))}" if division.get("head_name") else "")
        + f" · Дата анализа: {e(_date(division.get('analysis_completed_at')))}</div>"
    )
    kpis = [
        (stats["processes"], "процессов выявлено"),
        (f"{stats['potential']}%" if stats["potential"] is not None else "—", "средний потенциал автоматизации"),
        (stats["recommendations"], "рекомендаций"),
        (stats["fte"], "FTE — потенциальная экономия"),
        (stats["tobe"], "TO-BE моделей"),
        (stats["workflows"], "n8n workflows"),
    ]
    html.append("<div class='kpis'>" + "".join(f"<div class='kpi'><b>{e(v)}</b><span>{e(t)}</span></div>" for v, t in kpis) + "</div>")
    for warning in division.get("analysis_warnings") or []:
        html.append(f"<div class='warn'>⚠ {e(warning)}</div>")

    if data["scoring"]:
        html.append("<h3>Потенциал автоматизации по процессам</h3>")
        html.append(_svg_hbar(
            [(s.get("process_name", ""), s.get("automation_potential") or 0) for s in data["scoring"][:12]], 100, "%"
        ))
    if data["recommendations"]:
        colors = {"simple_automation": "#10b981", "classic_digitization": "#3b82f6", "ai_reengineering": "#8b5cf6"}
        html.append("<h3>Рекомендации по категориям</h3>")
        html.append(_svg_donut([
            (label, sum(1 for r in data["recommendations"] if r.get("category") == key), colors[key])
            for key, label in CATEGORY_LABELS.items()
        ]))

    html.append("<h2>1. Анкета подразделения</h2>")
    answers = (data.get("questionnaire") or {}).get("answers") or {}
    if answers:
        html.append("<table><tr><th style='width:38%'>Вопрос</th><th>Ответ</th></tr>")
        html.extend(f"<tr><td>{e(QUESTION_LABELS.get(k, k))}</td><td>{e(_answer(v))}</td></tr>" for k, v in answers.items())
        html.append("</table>")
    else:
        html.append("<p class='muted'>Анкета не заполнена</p>")

    html.append("<h2>2. Выявленные бизнес-процессы</h2>")
    if not data["processes"]:
        html.append("<p class='muted'>Процессы не выявлены</p>")
    for index, process in enumerate(data["processes"], 1):
        html.append(f"<div class='card'><h3>2.{index}. {e(process.get('name'))}</h3>")
        if process.get("description"):
            html.append(f"<p>{e(process['description'])}</p>")
        meta = [(label, process.get(key)) for label, key in (
            ("Владелец", "owner"), ("Триггер", "trigger"), ("Частота", "frequency"), ("SLA", "sla"))]
        html.append("".join(f"<div><b>{label}:</b> {e(value)}</div>" for label, value in meta if value))
        if process.get("roles"):
            html.append("<div><b>Роли:</b> " + "".join(f"<span class='tag'>{e(r)}</span>" for r in process["roles"]) + "</div>")
        if process.get("systems"):
            html.append("<div><b>Системы:</b> " + "".join(f"<span class='tag'>{e(s)}</span>" for s in process["systems"]) + "</div>")
        steps = process.get("steps_detail") or [{"name": s} for s in process.get("steps") or []]
        if steps:
            html.append("<ol>")
            for step in steps:
                extra = ", ".join(x for x in (step.get("role"), step.get("system")) if x)
                marker = "◇ " if step.get("type") == "decision" else ""
                html.append(f"<li>{marker}{e(step['name'])}" + (f" <span class='muted'>({e(extra)})</span>" if extra else "") + "</li>")
            html.append("</ol>")
        if process.get("pain_points"):
            html.append("<div><b>Проблемы:</b><ul>" + "".join(f"<li>{e(p)}</li>" for p in process["pain_points"]) + "</ul></div>")
        html.append("</div>")

    html.append("<h2>3. Скоринг процессов</h2>")
    if data["scoring"]:
        html.append("<table><tr><th>Процесс</th>" + "".join(f"<th>{e(SCORE_LABELS[f])}</th>" for f in SCORE_FIELDS) + "<th>Потенциал</th></tr>")
        for score in data["scoring"]:
            cells = "".join(
                f"<td class='num' style='background:#{_score_color(score.get(f) or 1)}'>{e(score.get(f, '—'))}</td>"
                for f in SCORE_FIELDS
            )
            html.append(f"<tr><td>{e(score.get('process_name'))}</td>{cells}<td class='num'><b>{e(score.get('automation_potential', '—'))}%</b></td></tr>")
        html.append("</table><p class='muted'>Шкала 1–10. Потенциал автоматизации — интегральный индекс 0–100%.</p>")
    else:
        html.append("<p class='muted'>Скоринг не выполнен</p>")

    html.append("<h2>4. Рекомендации</h2>")
    if not data["recommendations"]:
        html.append("<p class='muted'>Рекомендации не сформированы</p>")
    for key, label in CATEGORY_LABELS.items():
        items = [r for r in data["recommendations"] if r.get("category") == key]
        if not items:
            continue
        html.append(f"<h3>{e(label)}</h3>")
        for recommendation in items:
            html.append(
                f"<div class='card'><b>{e(recommendation.get('title'))}</b> "
                f"<span class='tag'>приоритет {e(recommendation.get('priority', '—'))}</span>"
                f"<span class='tag'>сложность: {e(COMPLEXITY_RU.get(recommendation.get('complexity'), '—'))}</span>"
                f"<span class='tag'>FTE {e(recommendation.get('estimated_fte_savings', 0))}</span>"
                + (f"<span class='tag'>{e(recommendation['implementation_weeks'])} нед.</span>" if recommendation.get("implementation_weeks") else "")
                + f"<p>{e(recommendation.get('description'))}</p>"
                + (f"<div><b>Эффект:</b> {e(recommendation['effect'])}</div>" if recommendation.get("effect") else "")
                + (f"<div><b>Стек:</b> {e(recommendation['implementation_stack'])}</div>" if recommendation.get("implementation_stack") else "")
                + "</div>"
            )

    html.append("<h2>5. TO-BE модели</h2>")
    if not data["tobe"]:
        html.append("<p class='muted'>TO-BE модели не сформированы</p>")
    for model in data["tobe"]:
        metrics = model.get("metrics") or {}
        html.append(f"<div class='card'><h3>{e(model.get('process_name'))}</h3>")
        if metrics.get("steps_before") is not None:
            html.append(
                f"<div><b>Шагов:</b> {e(metrics.get('steps_before'))} → {e(metrics.get('steps_after'))}, "
                f"автоматизировано: {e(metrics.get('automated_count', 0))}"
                + (f", сокращение времени ≈ {e(model['time_reduction_percent'])}%" if model.get("time_reduction_percent") else "")
                + "</div>"
            )
        if model.get("improvements"):
            html.append("<b>Улучшения:</b><ul>" + "".join(f"<li>{e(i)}</li>" for i in model["improvements"]) + "</ul>")
        new_steps = (model.get("new_process") or {}).get("steps_detail") or []
        if new_steps:
            html.append("<b>Целевой процесс:</b><ol>" + "".join(
                f"<li>{e(s['name'])}{' <span class=tag>авто</span>' if s.get('automated') else ''}</li>" for s in new_steps
            ) + "</ol>")
        if model.get("expected_effect"):
            html.append(f"<div><b>Эффект:</b> {e(model['expected_effect'])}</div>")
        if model.get("risks"):
            html.append(f"<div><b>Риски:</b> {e('; '.join(model['risks']))}</div>")
        html.append("</div>")

    html.append("<h2>6. n8n workflows</h2>")
    if not data["n8n_workflows"]:
        html.append("<p class='muted'>Workflows не сгенерированы</p>")
    for workflow in data["n8n_workflows"]:
        html.append(
            f"<div class='card'><b>{e(workflow.get('name'))}</b>"
            + (f"<p>{e(workflow['description'])}</p>" if workflow.get("description") else "")
            + (f"<div><b>Экономия:</b> {e(workflow['estimated_time_savings'])}</div>" if workflow.get("estimated_time_savings") else "")
            + (f"<div><b>Интеграции:</b> {e(', '.join(workflow['integration_points']))}</div>" if workflow.get("integration_points") else "")
            + "<div class='chain'>" + " → ".join(e(n.get("name", "")) for n in workflow.get("nodes") or []) + "</div></div>"
        )
    html.append("</body></html>")

    # Внешние ресурсы запрещены: в отчёт попадает текст пользователей и модели
    fetcher = URLFetcher(allowed_protocols={"data"})
    return HTML(string="".join(html), url_fetcher=fetcher).write_pdf()
