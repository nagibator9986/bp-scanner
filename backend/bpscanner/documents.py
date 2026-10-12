"""Извлечение текста из загруженных документов."""
import logging
from pathlib import Path

logger = logging.getLogger("bpscanner.documents")

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".csv", ".md"}
# Сколько извлечённого текста храним в базе на один документ
STORED_TEXT_LIMIT = 300_000


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # битые страницы не должны ронять весь документ
            logger.warning("PDF page extraction failed in %s: %s", path.name, exc)
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and (not cells or cells[-1] != text):  # объединённые ячейки повторяются
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"# Лист: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(c) for c in row if c is not None and str(c).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"# Слайд {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Заметки: {notes}")
    return "\n".join(parts)


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".txt": _read_text_file,
    ".csv": _read_text_file,
    ".md": _read_text_file,
}


def extract_text_from_file(file_path: str | Path) -> str:
    """Синхронное извлечение текста. Вызывайте через asyncio.to_thread."""
    path = Path(file_path)
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None or not path.is_file():
        return ""
    try:
        text = extractor(path)
    except Exception as exc:
        logger.error("Error extracting text from %s: %s", path.name, exc)
        return ""
    # Схлопываем пустые строки, чтобы не тратить бюджет промпта
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line.strip())


def build_documents_context(documents: list[tuple[str, str]], budget: int) -> str:
    """Склеивает тексты документов, честно деля общий бюджет символов между ними.

    Короткие документы берутся целиком, а освободившийся бюджет достаётся длинным.
    """
    documents = [(name, text) for name, text in documents if text and text.strip()]
    if not documents:
        return "Документы не загружены"

    remaining = budget
    allocation: dict[int, int] = {}
    pending = sorted(range(len(documents)), key=lambda i: len(documents[i][1]))
    while pending:
        share = remaining // len(pending)
        index = pending.pop(0)
        take = min(len(documents[index][1]), share)
        allocation[index] = take
        remaining -= take

    parts = []
    for index, (name, text) in enumerate(documents):
        limit = allocation.get(index, 0)
        chunk = text[:limit]
        suffix = "\n[...текст сокращён...]" if limit < len(text) else ""
        parts.append(f"=== Документ: {name} ===\n{chunk}{suffix}")
    return "\n\n".join(parts)
